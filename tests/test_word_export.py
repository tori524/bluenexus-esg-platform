from io import BytesIO
from docx import Document
from calculations import calculate
from report_export import build_esg_report_docx

def test_word_export_contains_charts_and_disclosure_tables():
    current={'influent_m3':1000,'treated_m3':750,'electricity_kwh':800,'grid_ef_kgco2e_kwh':0.4}
    previous={'influent_m3':1000,'treated_m3':700,'electricity_kwh':1000}
    values={**current,'baseline_treated_m3':700,'baseline_electricity_kwh':1000}
    b=build_esg_report_docx('Demo','2026',{'results':calculate(values),'current':current,'previous':previous,'estimated':True,'ef':0.4})
    doc=Document(BytesIO(b))
    text=' '.join(p.text for p in doc.paragraphs)
    assert 'Illustrative / estimated' in text
    assert 'GRI / AWS reporting support' in text
    assert len(doc.inline_shapes)>=3
    assert len(doc.tables)==2
